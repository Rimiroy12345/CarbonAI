import io
import sys
import unittest
from pathlib import Path
from openpyxl import load_workbook
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from reports import build_workbook

class ReportTests(unittest.TestCase):
    def test_formulas_units_and_injection_protection(self):
        data={'company_name':'=HYPERLINK("https://example.com")','company_id':'test', 'calculated_at':'2026-10-11','total_tco2e':2,
            'activities':[{'category':'electricity','quantity':1000,'unit':'kWh','scope':2,'factor':0.7}]}
        workbook=load_workbook(io.BytesIO(build_workbook(data)))
        self.assertEqual(workbook['Summary']['B1'].data_type,'s')
        self.assertEqual(workbook['Activities']['F2'].value,'=ROUND(B2*E2,2)')
        self.assertEqual(workbook['Activities']['C2'].value,'kWh')
        self.assertEqual(workbook['Summary']['B12'].value,'=SUM(B9:B11)')
        self.assertEqual(workbook.sheetnames,['Summary','Activities','Method'])

if __name__ == '__main__': unittest.main()
