"""Runtime Excel exports using the application's existing workbook dependency."""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

def build_workbook(data):
    workbook = Workbook()
    summary = workbook.active
    summary.title = 'Summary'
    summary.append(['CarbonAI assessment', data['company_name']])
    for label, key in [('Assessment ID', 'company_id'), ('Saved date (UTC)', 'calculated_at'), ('Industry', 'industry'), ('Location', 'location'), ('Employees', 'employee_count'), ('Saved total tCO2e', 'total_tco2e')]:
        summary.append([label, data.get(key)])
    summary.append(['Scope', 'Current-factor tCO2e'])
    activities = workbook.create_sheet('Activities')
    activities.append(['Activity', 'Quantity', 'Unit', 'Scope', 'Factor kgCO2e per unit', 'Calculated kgCO2e'])
    for index, row in enumerate(data['activities'], 2):
        activities.append([row['category'], row['quantity'], row['unit'], row['scope'], row['factor'], f'=ROUND(B{index}*E{index},2)' if row['factor'] is not None else 'Missing factor'])
    last = len(data['activities']) + 1
    for scope in [1, 2, 3]:
        summary.append([f'Scope {scope}', f'=SUMIF(Activities!D2:D{last},{scope},Activities!F2:F{last})/1000'])
    summary.append(['Current-factor total tCO2e', '=SUM(B9:B11)'])
    notes = workbook.create_sheet('Method')
    notes.append(['Method and limitations'])
    for note in ['Inputs: user-entered or reviewed annual activity totals for the selected assessment.',
        'Factors: current CarbonAI database factors; historical factor snapshots are not retained.',
        'Saved total may differ from current-factor recalculation. Formula cells recalculate when opened in Excel.',
        'Scope 1 assumes company-owned or controlled natural gas, petrol and diesel.',
        'Scope 2 includes purchased electricity; no market-based claim.',
        'Scope 3 covers reported air travel, hotels, commuting and waste only, not a complete inventory.',
        'Missing factors require review; do not treat incomplete calculations as zero emissions.',
        'Original uploaded documents and extraction evidence are not retained in saved assessments.',
        'AI advice is in the PDF report. This workbook contains no ML forecast.']:
        notes.append([note])
    workbook.calculation.fullCalcOnLoad = True
    for sheet in workbook:
        sheet.freeze_panes = 'A2'
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='163F32')
        for row in sheet:
            for cell in row:
                # Prevent company/user strings from becoming spreadsheet formulas.
                if cell.data_type == 'f' and not (sheet.title == 'Activities' and cell.column == 6 or sheet.title == 'Summary' and cell.row >= 9 and cell.column == 2):
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if isinstance(cell.value, (int, float)) or cell.data_type == 'f':
                    cell.number_format = '#,##0.0000'
        for column in ['A', 'B', 'C', 'D', 'E', 'F']:
            sheet.column_dimensions[column].width = 28 if column in ['A', 'E', 'F'] else 22
    notes.column_dimensions['A'].width = 110
    for row in notes:
        notes.row_dimensions[row[0].row].height = 32
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()
